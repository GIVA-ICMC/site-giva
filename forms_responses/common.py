#!/usr/bin/env python3
"""Shared helpers for GIVA Excel importers."""

from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
FORMS_DIR = Path(__file__).resolve().parent
STUDENT_XLSX = FORMS_DIR / "Informações para o site (respostas).xlsx"
FACULTY_XLSX = FORMS_DIR / "Dados site GIVA (docentes).xlsx"
PHOTO_DIR = FORMS_DIR / "Informações para o site (File responses)"
PEOPLE_DIR = REPO_ROOT / "content" / "people"
PUBS_DIR = REPO_ROOT / "content" / "publication"
POSTS_DIR = REPO_ROOT / "content" / "post"
UPLOADS_DIR = REPO_ROOT / "static" / "uploads"

STUDENT_SHEET = "Respostas ao formulário 1"

COL_FULL_NAME = "Nome Inteiro (Full Name)"
COL_EMAIL = "Email USP (or other if preferred, or not available)"
COL_DEGREES = "Títulos e onde os obteu (Degrees and institutions)"
COL_BIO = "Your description. In english!"
COL_PHOTO = "Foto de perfil (Profile Picture) optional, max 1MB"
COL_OTHERS = "Others"
COL_LEVEL = "Nível de pesquisador (Researcher level)"
COL_INTERESTS = "Tópicos de pesquisa / Interesse (Topic of research / Interests)"
COL_PAPERS = "Lista de artigos publicado e informações (List of Published Papers and Links)"

SOCIAL_COLS: list[tuple[str, str]] = [
    (COL_EMAIL, "envelope"),
    ("Github", "github"),
    ("Linkedin", "linkedin"),
    ("Google Scholar", "google-scholar"),
    ("Youtube", "youtube"),
    ("Instagram", "instagram"),
    ("Twitter", "twitter"),
    ("Facebook", "facebook"),
    ("Curriculum Lattes", "lattes"),
]

LEVEL_MAP: dict[str, tuple[str, str]] = {
    "phd": ("Grad Students", "PhD Candidate"),
    "master's": ("Grad Students", "Master's Student"),
    "ic (junior researcher)": ("Undergraduate Students", "Undergraduate Researcher (IC)"),
    "post-doc": ("Researchers", "Postdoctoral Researcher"),
    "direct phd (no master's)": ("Grad Students", "PhD Candidate"),
}

_DEGREE_KEYWORDS = (
    r"Gradua", r"Bacharel", r"Bachelor", r"Licencia", r"Mestr", r"Master",
    r"Doutor", r"PhD", r"Post", r"P\u00f3s", r"Especiali", r"T\u00e9cnico",
    r"Atualmente", r"Currently",
)
_DEGREE_START = re.compile(
    r",\s*(?=(?:" + "|".join(_DEGREE_KEYWORDS) + r"))",
    re.IGNORECASE,
)
_CURRENT_RE = re.compile(r"^(Atualmente|Currently)\b", re.IGNORECASE)

DOI_RE = re.compile(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)
ARXIV_RE = re.compile(
    r"(?:arxiv\.org/(?:abs|pdf)/|arxiv:)(\d{4}\.\d{4,5}(?:v\d+)?)",
    re.IGNORECASE,
)
URL_RE = re.compile(r"https?://[^\s,;|<>]+", re.IGNORECASE)
ORCID_RE = re.compile(r"orcid\.org", re.IGNORECASE)
GITHUB_RE = re.compile(r"github\.com", re.IGNORECASE)
YEAR_RE = re.compile(r"\b(20\d{2}|19\d{2})\b")

PHOTO_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}

ICON_BY_HOST = [
    (re.compile(r"orcid\.org", re.I), "orcid"),
    (re.compile(r"github\.com", re.I), "github"),
    (re.compile(r"linkedin\.com", re.I), "linkedin"),
    (re.compile(r"scholar\.google", re.I), "google-scholar"),
    (re.compile(r"lattes\.cnpq|buscatextual\.cnpq", re.I), "lattes"),
    (re.compile(r"researchgate\.net", re.I), "researchgate"),
    (re.compile(r"youtube\.com|youtu\.be", re.I), "youtube"),
    (re.compile(r"instagram\.com", re.I), "instagram"),
    (re.compile(r"twitter\.com|\bx\.com/", re.I), "twitter"),
    (re.compile(r"facebook\.com", re.I), "facebook"),
]


def slugify(text: str) -> str:
    """Convert a name or title to a URL-safe ASCII slug."""
    text = unicodedata.normalize("NFD", str(text))
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s-]", "", text)
    text = re.sub(r"\s+", "-", text.strip())
    text = re.sub(r"-+", "-", text)
    return text.strip("-")


def qs(s: str) -> str:
    """Wrap a value in a YAML double-quoted string."""
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def cell_str(value) -> str:
    """Normalise an Excel cell value to a stripped string."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_sheet_rows(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, str]]]:
    """Return (headers, dict rows) from the first/named sheet of an xlsx file."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return [], []
    headers = [cell_str(h) for h in header_row]
    out: list[dict[str, str]] = []
    for raw in rows_iter:
        row = {headers[i]: cell_str(raw[i]) if i < len(raw) else "" for i in range(len(headers))}
        if any(row.values()):
            out.append(row)
    wb.close()
    return headers, out


def faculty_sheet_rows(ws) -> list[list[str]]:
    """Return all rows of a faculty sheet as lists of strings."""
    rows = []
    for raw in ws.iter_rows(values_only=True):
        rows.append([cell_str(c) for c in raw])
    return rows


def load_faculty_sheets(path: Path) -> list[tuple[str, list[list[str]]]]:
    """Load every faculty tab as (sheet_name, rows)."""
    wb = load_workbook(path, data_only=True)
    try:
        return [(name, faculty_sheet_rows(wb[name])) for name in wb.sheetnames]
    finally:
        wb.close()


def faculty_value(rows: list[list[str]], label_prefix: str, col: int = 1) -> str:
    """Find a row whose first cell starts with label_prefix and return column `col`."""
    prefix = label_prefix.lower()
    for row in rows:
        if row and row[0].lower().startswith(prefix):
            return row[col] if col < len(row) else ""
    return ""


def faculty_labeled_rows(rows: list[list[str]], pattern: str) -> list[list[str]]:
    """Rows whose first cell matches a regex (e.g. r'Projeto \\d+')."""
    cre = re.compile(pattern, re.IGNORECASE)
    return [row for row in rows if row and cre.fullmatch(row[0].strip())]


def maybe_write(path: Path, content: str, overwrite: bool) -> str:
    """Write `content` to `path`. Returns 'create', 'overwrite', or 'skip'."""
    existed = path.exists()
    if existed and not overwrite:
        return "skip"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    return "overwrite" if existed else "create"


def parse_interests(raw: str) -> list[str]:
    """Prefer English text in parentheses; otherwise keep the token."""
    if not raw.strip():
        return []
    items = []
    for part in re.split(r"[,;]", raw):
        part = part.strip()
        if not part:
            continue
        m = re.search(r"\(([^)]+)\)", part)
        items.append(m.group(1).strip() if m else part)
    return items


def parse_degrees(raw: str) -> list[dict]:
    """Parse the free-text degrees field into {course, institution?, year?} dicts."""
    if not raw.strip():
        return []

    if "\n" in raw or re.search(r"\d+\)\s", raw):
        lines = re.split(r"\n|(?=\d+\)\s)", raw)
    else:
        lines = _DEGREE_START.split(raw)

    courses = []
    for line in lines:
        line = re.sub(r"^\d+\)\s*", "", line).strip().rstrip(";")
        if not line or _CURRENT_RE.match(line):
            continue
        year_m = YEAR_RE.search(line)
        year = int(year_m.group(1)) if year_m else None

        sep = re.split(
            r"\s+[-–]\s+|\s+pela\s+|\s+pelo\s+|\s+no\s+|\s+na\s+|\s+at\s+|\s+by\s+",
            line, maxsplit=1, flags=re.IGNORECASE,
        )
        if len(sep) == 2:
            course_part = sep[0].strip()
            inst_part = sep[1].strip()
            inst_part = re.sub(r"\s*\(.*?\)\s*", "", inst_part)
            inst_part = re.sub(r"\b(20\d{2}|19\d{2})\b.*", "", inst_part).strip().rstrip(".,")
        else:
            course_part = line
            inst_part = ""

        course_part = re.sub(r"\s*[\(\[]?\b(20\d{2}|19\d{2})\b[\)\]]?.*", "", course_part).strip()
        entry: dict = {"course": course_part}
        if inst_part:
            entry["institution"] = inst_part
        if year:
            entry["year"] = year
        courses.append(entry)
    return courses


def parse_social(row: dict) -> list[dict]:
    """Build the social list from student-form columns."""
    links = []
    for col, icon in SOCIAL_COLS:
        val = row.get(col, "").strip()
        if not val:
            continue
        if icon == "envelope":
            if "@" not in val:
                continue
            href = f"mailto:{val}"
        elif icon == "instagram" and val.startswith("@"):
            href = f"https://instagram.com/{val.lstrip('@')}"
        else:
            href = val
        links.append({"icon": icon, "link": href})

    others = row.get(COL_OTHERS, "").strip()
    if others:
        for token in re.split(r"[\n,]", others):
            token = token.strip()
            if not token:
                continue
            url_m = URL_RE.search(token)
            url = url_m.group(0).rstrip(").,;") if url_m else ""
            if re.search(r"orcid\.org", token, re.I):
                links.append({"icon": "orcid", "link": url or token})
            elif re.search(r"researchgate\.net", token, re.I):
                links.append({"icon": "researchgate", "link": url or token})
    return links


def parse_links_blob(raw: str) -> list[dict]:
    """Parse a free-text blob of labelled URLs into social icon/link dicts."""
    if not raw.strip():
        return []
    links = []
    seen = set()
    for token in re.split(r"[\n|]", raw):
        token = token.strip()
        if not token:
            continue
        for url in URL_RE.findall(token):
            url = url.rstrip(").,;")
            if url in seen:
                continue
            seen.add(url)
            icon = "link"
            for cre, name in ICON_BY_HOST:
                if cre.search(url):
                    icon = name
                    break
            links.append({"icon": icon, "link": url})
    return links


def split_name(full_name: str) -> tuple[str, str]:
    parts = full_name.split()
    first = parts[0] if parts else full_name
    last = " ".join(parts[1:]) if len(parts) > 1 else ""
    return first, last


def gdrive_file_id(url: str) -> str | None:
    m = re.search(r"[?&/]id=([A-Za-z0-9_-]+)", url)
    if m:
        return m.group(1)
    m = re.search(r"/file/d/([A-Za-z0-9_-]+)", url)
    return m.group(1) if m else None


def photo_matches(filename: str, full_name: str) -> bool:
    """True if a File-responses photo belongs to this person."""
    stem = Path(filename).stem
    suffix = stem.rsplit(" - ", 1)[-1]
    name_slug = slugify(full_name)
    suffix_slug = slugify(suffix)
    if not suffix_slug or not name_slug:
        return False
    if suffix_slug == name_slug or suffix_slug in name_slug or name_slug in suffix_slug:
        return True
    suffix_parts = [p for p in suffix_slug.split("-") if p not in {"de", "da", "do", "dos", "das"}]
    name_parts = set(name_slug.split("-"))
    return bool(suffix_parts) and all(p in name_parts for p in suffix_parts)


def find_photo(full_name: str) -> Path | None:
    if not PHOTO_DIR.exists():
        return None
    matches = [
        p for p in PHOTO_DIR.rglob("*")
        if p.is_file() and p.suffix.lower() in PHOTO_EXTS and photo_matches(p.name, full_name)
    ]
    if not matches:
        return None
    return max(matches, key=lambda p: p.stat().st_size)


def copy_photo(src: Path, slug: str) -> Path:
    """Copy a profile photo into static/uploads/<slug><ext>. Returns the destination."""
    dest = UPLOADS_DIR / f"{slug}{src.suffix.lower()}"
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    if not dest.exists() and dest.resolve() != src.resolve():
        dest.write_bytes(src.read_bytes())
    return dest


def extract_doi(text: str) -> str | None:
    m = DOI_RE.search(text or "")
    if not m:
        return None
    doi = m.group(0).rstrip(").,;")
    doi = re.sub(r"/(?:full|pdf|html|meta)$", "", doi, flags=re.I)
    return doi


def extract_arxiv_id(text: str) -> str | None:
    m = ARXIV_RE.search(text or "")
    return m.group(1) if m else None


def looks_like_url(text: str) -> bool:
    return bool(URL_RE.search(text or ""))


def _split_commas(line: str) -> list[str]:
    parts = [p.strip() for p in re.split(r",\s+(?=https?://|10\.\d{4})", line) if p.strip()]
    if len(parts) == 1 and "," in line and looks_like_url(line):
        parts = [p.strip() for p in line.split(",") if p.strip()]
    return parts or [line]


def split_paper_tokens(raw: str) -> list[str]:
    """Split a mixed papers cell into individual tokens (URLs, DOIs, titles).

    `|` separates multiple papers when any segment is a URL/DOI; otherwise it
    is treated as `title | venue` and only the title is kept.
    """
    if not raw.strip():
        return []
    chunks: list[str] = []
    for line in re.split(r"\n+", raw):
        line = line.strip()
        if not line:
            continue
        if "|" in line:
            segs = [p.strip() for p in line.split("|") if p.strip()]
            if any(extract_doi(p) or looks_like_url(p) or extract_arxiv_id(p) for p in segs):
                for p in segs:
                    chunks.extend(_split_commas(p))
            else:
                chunks.append(segs[0])
        else:
            chunks.extend(_split_commas(line))
    return chunks


def latest_year(text: str) -> int | None:
    years = [int(y) for y in YEAR_RE.findall(text or "")]
    return max(years) if years else None


def load_people_index() -> list[dict]:
    """Index existing people pages for author matching."""
    people = []
    if not PEOPLE_DIR.exists():
        return people
    for path in PEOPLE_DIR.glob("*.md"):
        if path.name.startswith("_"):
            continue
        text = path.read_text(encoding="utf-8")
        title_m = re.search(r'^title:\s*"?([^"\n]+)"?\s*$', text, re.M)
        first_m = re.search(r'^first_name:\s*"?([^"\n]+)"?\s*$', text, re.M)
        last_m = re.search(r'^last_name:\s*"?([^"\n]+)"?\s*$', text, re.M)
        title = title_m.group(1).strip() if title_m else path.stem.replace("-", " ").title()
        people.append({
            "slug": path.stem,
            "title": title,
            "first": (first_m.group(1).strip() if first_m else "").strip('"'),
            "last": (last_m.group(1).strip() if last_m else "").strip('"'),
            "title_slug": slugify(title),
        })
    return people


def match_author(name: str, people: list[dict]) -> str:
    """Return a people slug when the name clearly matches; otherwise the original name."""
    nslug = slugify(name)
    if not nslug:
        return name
    for p in people:
        if nslug in {p["slug"], p["title_slug"]}:
            return p["slug"]
    name_parts = [w for w in nslug.split("-") if w not in {"de", "da", "do", "dos", "das", "e"}]
    if not name_parts:
        return name
    last = name_parts[-1]
    first = name_parts[0]
    candidates = []
    for p in people:
        p_last = slugify(p["last"]).split("-")[-1] if p["last"] else p["slug"].split("-")[-1]
        p_first = slugify(p["first"]) if p["first"] else p["slug"].split("-")[0]
        if last == p_last and (p_first.startswith(first) or first.startswith(p_first[:3])):
            candidates.append(p["slug"])
    if len(candidates) == 1:
        return candidates[0]
    return name


def render_person_md(
    *,
    full_name: str,
    group: str,
    role: str,
    interests: list[str],
    degrees: list[dict],
    socials: list[dict],
    description: str,
    avatar_rel: str | None,
    avatar_comment: str | None = None,
) -> str:
    first, last = split_name(full_name)
    fm = ["---"]
    fm.append(f"title: {qs(full_name)}")
    fm.append(f"first_name: {qs(first)}")
    fm.append(f"last_name: {qs(last)}")
    fm.append(f'userGroup: ["{group}"]')
    fm.append(f"role: {qs(role)}")
    fm.append("organizations:")
    fm.append('  - name: "ICMC-USP"')
    fm.append('    url: "https://www.icmc.usp.br/"')
    if socials:
        fm.append("social:")
        for s in socials:
            fm.append(f'  - icon: {qs(s["icon"])}')
            fm.append(f'    link: {qs(s["link"])}')
    if interests:
        fm.append("interests:")
        for i in interests:
            fm.append(f"  - {qs(i)}")
    if degrees:
        fm.append("education:")
        fm.append("  courses:")
        for d in degrees:
            fm.append(f'    - course: {qs(d["course"])}')
            if d.get("institution"):
                fm.append(f'      institution: {qs(d["institution"])}')
            if d.get("year"):
                fm.append(f"      year: {d['year']}")
    if avatar_rel:
        fm.append(f"avatar: {qs(avatar_rel)}")
    if avatar_comment:
        fm.append(avatar_comment)
    fm.append("---")
    body = f"\n{description}\n" if description else "\n"
    return "\n".join(fm) + body
